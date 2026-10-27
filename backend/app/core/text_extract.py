"""Document text extraction for PDF, DOCX, TXT and scanned PDFs (OCR via the tesseract binary).

All uploaded content is treated as untrusted. Type is validated by magic bytes, not by the
file extension alone.
"""
from __future__ import annotations

import io
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import asdict, dataclass, field
from pathlib import Path

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class ExtractionError(Exception):
    pass


@dataclass
class ExtractedDocument:
    text: str
    pages: list[dict] = field(default_factory=list)  # [{page, start, end}]
    method: str = ""
    warnings: list[str] = field(default_factory=list)
    detected_type: str = ""

    def to_dict(self) -> dict:
        return asdict(self)


def detect_type(data: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ExtractionError(f"File type {ext or '(none)'} is not supported. Use PDF, DOCX or TXT.")
    if not data:
        raise ExtractionError("The file is empty.")
    if ext == ".pdf":
        if not data[:1024].lstrip().startswith(b"%PDF"):
            raise ExtractionError("The file has a .pdf extension but is not a PDF document.")
        return "pdf"
    if ext == ".docx":
        if not data.startswith(b"PK\x03\x04"):
            raise ExtractionError("The file has a .docx extension but is not a DOCX archive.")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = zf.namelist()
                if "word/document.xml" not in names:
                    raise ExtractionError("The DOCX archive has no word/document.xml part.")
                total = sum(i.file_size for i in zf.infolist())
                if total > 200 * 1024 * 1024:
                    raise ExtractionError("The DOCX archive expands beyond the safe size limit.")
        except zipfile.BadZipFile as exc:
            raise ExtractionError("The DOCX archive is corrupt.") from exc
        return "docx"
    if b"\x00" in data[:4096]:
        raise ExtractionError("The .txt file contains binary content.")
    return "txt"


def ocr_available() -> bool:
    return shutil.which("tesseract") is not None


def _ocr_pdf(data: bytes, max_pages: int) -> list[str]:
    import pypdfium2 as pdfium  # optional dependency, imported lazily

    texts = []
    pdf = pdfium.PdfDocument(data)
    with tempfile.TemporaryDirectory() as tmp:
        for i in range(min(len(pdf), max_pages)):
            img = pdf[i].render(scale=300 / 72).to_pil()
            path = Path(tmp) / f"p{i}.png"
            img.save(path)
            res = subprocess.run(["tesseract", str(path), "stdout"], capture_output=True, timeout=120, check=False)
            if res.returncode != 0:
                raise ExtractionError(f"OCR failed on page {i + 1}: {res.stderr.decode(errors='ignore')[:200]}")
            texts.append(res.stdout.decode("utf-8", errors="replace"))
    return texts


def _join_pages(page_texts: list[str]) -> tuple[str, list[dict]]:
    parts, pages, pos = [], [], 0
    for i, t in enumerate(page_texts, start=1):
        t = t.replace("\r\n", "\n").replace("\x00", "")
        parts.append(t)
        pages.append({"page": i, "start": pos, "end": pos + len(t)})
        pos += len(t) + 2
    return "\n\n".join(parts), pages


def extract_text(data: bytes, filename: str, max_bytes: int = 10 * 1024 * 1024, enable_ocr: bool = True,
                 max_pages: int = 30, min_chars_per_page: int = 25) -> ExtractedDocument:
    if len(data) > max_bytes:
        limit = f"{max_bytes / (1024 * 1024):.1f} MB" if max_bytes >= 1024 * 1024 else f"{max_bytes} bytes"
        raise ExtractionError(f"File exceeds the {limit} limit.")
    kind = detect_type(data, filename)
    warnings: list[str] = []
    if kind == "txt":
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
            warnings.append("Decoded as Latin-1 because the file is not valid UTF-8.")
        text = text.replace("\r\n", "\n")
        return ExtractedDocument(text, [{"page": 1, "start": 0, "end": len(text)}], "plain-text", warnings, kind)
    if kind == "docx":
        import docx  # python-docx

        try:
            d = docx.Document(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001
            raise ExtractionError(f"Could not read DOCX: {exc}") from exc
        lines = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                lines.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(lines)
        warnings.append("DOCX has no fixed pagination; offsets are relative to the whole document.")
        return ExtractedDocument(text, [{"page": 1, "start": 0, "end": len(text)}], "python-docx", warnings, kind)
    # PDF
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("Encrypted PDFs are not accepted.")
        if len(reader.pages) > max_pages:
            warnings.append(f"Only the first {max_pages} pages were processed.")
        page_texts = [(p.extract_text() or "") for p in reader.pages[:max_pages]]
    except ExtractionError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise ExtractionError(f"Could not parse PDF: {exc}") from exc
    sparse = sum(1 for t in page_texts if len(t.strip()) < min_chars_per_page)
    method = "pypdf"
    if page_texts and sparse == len(page_texts):
        if enable_ocr and ocr_available():
            page_texts = _ocr_pdf(data, max_pages)
            method = "tesseract-ocr"
            warnings.append("No text layer found; text was produced by OCR and may contain recognition errors.")
        else:
            warnings.append("No text layer found and OCR is not configured; the document needs OCR.")
    text, pages = _join_pages(page_texts)
    return ExtractedDocument(text, pages, method, warnings, kind)
