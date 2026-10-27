import io
import unittest

from app.core.text_extract import ExtractionError, extract_text, ocr_available


class TestExtraction(unittest.TestCase):
    def test_txt_and_type_checks(self):
        d = extract_text(b"Skills\nPython", "cv.txt")
        self.assertEqual(d.method, "plain-text")
        with self.assertRaises(ExtractionError):
            extract_text(b"not a pdf", "cv.pdf")
        with self.assertRaises(ExtractionError):
            extract_text(b"x", "cv.exe")
        with self.assertRaises(ExtractionError):
            extract_text(b"", "cv.txt")
        with self.assertRaises(ExtractionError):
            extract_text(b"a" * 2000, "cv.txt", max_bytes=1000)

    def test_docx(self):
        import docx
        doc = docx.Document()
        doc.add_paragraph("Skills")
        doc.add_paragraph("Python and Docker")
        buf = io.BytesIO()
        doc.save(buf)
        d = extract_text(buf.getvalue(), "cv.docx")
        self.assertIn("Python and Docker", d.text)

    def test_text_pdf(self):
        from reportlab.pdfgen import canvas
        buf = io.BytesIO()
        c = canvas.Canvas(buf)
        c.drawString(72, 720, "Skills: Python, PostgreSQL")
        c.showPage()
        c.save()
        d = extract_text(buf.getvalue(), "cv.pdf")
        self.assertEqual(d.method, "pypdf")
        self.assertIn("PostgreSQL", d.text)
        self.assertEqual(d.pages[0]["page"], 1)

    @unittest.skipUnless(ocr_available(), "tesseract not installed")
    def test_scanned_pdf_ocr(self):
        from PIL import Image, ImageDraw, ImageFont
        img = Image.new("RGB", (1700, 400), "white")
        draw = ImageDraw.Draw(img)
        try:
            font = ImageFont.truetype("DejaVuSans.ttf", 64)
        except OSError:
            font = ImageFont.load_default()
        draw.text((60, 120), "Skills Python Kubernetes", fill="black", font=font)
        buf = io.BytesIO()
        img.save(buf, format="PDF", resolution=200)
        d = extract_text(buf.getvalue(), "scan.pdf")
        self.assertEqual(d.method, "tesseract-ocr")
        self.assertIn("Kubernetes", d.text)
